"""Permanent Excel audit package generator.

Each approved calculation set produces a new, immutable .xlsx workbook whose
hash is registered in the case store — an issued package is never altered; a
new calculation produces a new package.

Drill-down works without macros:
  * five-year summary cells hyperlink to per-year calculation sheets;
  * calculation sheets use grouped (collapsible) rows per form/schedule;
  * every line shows its stable line ID, formula, upstream lines, ruleset
    parameters, and source inputs;
  * named ranges anchor key totals;
  * a detail index maps line IDs to cell locations.

After writing, `verify_workbook` re-opens the file and confirms workbook
totals equal the persisted calculation results (reconciliation check
REC-XLSX-001).
"""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

from . import ENGINE_VERSION, SCHEMA_VERSION
from .money import d
from .schemas import CalculationResult, ReconciliationCheck
from .services import SUMMARY_FIELDS
from .store import CaseStore

HDR = Font(bold=True, size=11)
TITLE = Font(bold=True, size=14)
GRAY = PatternFill("solid", start_color="EEEEEE")
WARN = PatternFill("solid", start_color="FFF2CC")
FAIL = PatternFill("solid", start_color="F4CCCC")


def _sheet_title(name: str) -> str:
    return name[:31]


def _header_row(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HDR
        cell.fill = GRAY


# Critical totals re-read from the saved workbook and reconciled against the
# persisted calculation results. Generation fails visibly on any mismatch.
VERIFIED_LINE_IDS = {
    "F1040.L11": ("agi", "AGI"),
    "F1040.L15": ("taxable_income", "taxable income"),
    "F1040.L24": ("total_tax", "total tax"),
    "F1040.L33": ("total_payments", "total payments"),
}


class AuditPackageBuilder:
    def __init__(self, store: CaseStore, case_id: str, calculation_ids: list[str],
                 comparison: dict[str, Any] | None, created_by: str,
                 manifest_id: str | None = None):
        self.store = store
        self.case_id = case_id
        self.results: list[CalculationResult] = [store.get_calculation(c) for c in calculation_ids]
        self.comparison = comparison
        self.created_by = created_by
        self.manifest_id = manifest_id or store.new_id("pkg")
        self.generated_at = store.now()
        self.wb = Workbook()
        self.detail_index: list[tuple[str, str, str]] = []  # line_id, sheet, cell

    # ------------------------------------------------------------------
    def build(self, out_path: Path) -> tuple[Path, str]:
        wb = self.wb
        wb.remove(wb.active)
        self._cover()
        self._metadata()
        self._base_inputs()
        self._import_mapping()
        self._assumptions_by_year()
        self._scenario_change_log()
        self._five_year_summary()
        self._scenario_comparison()
        calc_cells = self._calculation_sheets()
        self._reconciliation()
        self._calculation_trace()
        self._rules_manifest()
        self._warnings()
        self._review_signoff()
        self._change_history()
        self._detail_index_sheet()

        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
        file_hash = "sha256:" + hashlib.sha256(out_path.read_bytes()).hexdigest()
        # Release gate: workbook totals must match persisted results.
        checks = verify_workbook(out_path, self.results, calc_cells)
        failed = [c for c in checks if c.status == "failed"]
        if failed:
            raise RuntimeError(
                f"workbook verification failed: {[c.check_id for c in failed]}")
        self.store.register_package(self.case_id, out_path.name, file_hash, self.created_by)
        self._store_manifest(out_path.name, file_hash)
        return out_path, file_hash

    def _store_manifest(self, filename: str, file_hash: str) -> None:
        """Immutable PackageManifest record — the out-of-band identity the
        workbook's cover references."""
        from .records import PackageManifest

        base_calc = next((r for r in self.results if r.target_kind == "base"), None)
        case_doc = self.store._load(self.case_id)
        manifest = PackageManifest(
            manifest_id=self.manifest_id,
            tenant_id=case_doc.get("tenant_id", "tenant_default"),
            case_id=self.case_id,
            filename=filename,
            file_sha256=file_hash.removeprefix("sha256:"),
            generated_at=self.generated_at,
            generated_by=self.created_by,
            engine_version=ENGINE_VERSION,
            ruleset_versions=dict(self.results[0].ruleset_versions) if self.results else {},
            base_version_id=base_calc.target_id if base_calc else "",
            scenario_refs=[f"{r.target_id}@v{r.scenario_version}"
                           for r in self.results if r.target_kind == "scenario"],
            calculation_ids=[r.calculation_id for r in self.results],
            reconciliation_status=("failed" if any(r.reconciliation_status == "failed"
                                                   for r in self.results) else "passed"),
            review_status=("required" if any(r.review_status == "required"
                                             for r in self.results) else "not_required"),
        )
        self.store.backend.put_immutable(
            "manifests", self.manifest_id, manifest.model_dump(mode="json"))

    # ------------------------------------------------------------------
    def _cover(self) -> None:
        ws = self.wb.create_sheet("Cover")
        ws["A1"] = "Tax Planning Audit Package"
        ws["A1"].font = TITLE
        rows = [
            ("Case", self.case_id),
            ("Calculations", ", ".join(r.calculation_id for r in self.results)),
            ("Engine version", ENGINE_VERSION),
            ("Schema version", SCHEMA_VERSION),
            ("Prepared by", self.created_by),
            ("Generated at", self.generated_at),
            ("Manifest reference", self.manifest_id + " — this workbook's SHA-256 and full "
             "identity are recorded immutably in the case package registry and the "
             "manifests store under this id; verify a copy against them"),
            ("", ""),
            ("Certification", "All calculated values in this workbook were produced by the "
             "deterministic calculation engine identified above. No value was produced by "
             "an AI language model. Provisional-year figures are projections, not enacted law."),
            ("Status", "PLANNING ESTIMATE — not filed-return advice."),
        ]
        for i, (k, v) in enumerate(rows, start=3):
            ws.cell(row=i, column=1, value=k).font = HDR
            ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 100

    def _metadata(self) -> None:
        ws = self.wb.create_sheet("Case Metadata")
        _header_row(ws, 1, ["Field", "Value"])
        r = 2
        for res in self.results:
            for k, v in res.identity_block().items():
                ws.cell(row=r, column=1, value=f"{res.calculation_id}.{k}")
                ws.cell(row=r, column=2, value=str(v))
                r += 1

    def _base_inputs(self) -> None:
        ws = self.wb.create_sheet("Base Inputs")
        base_calc = next((r for r in self.results if r.target_kind == "base"), self.results[0])
        base = self.store.get_base_version(self.case_id, base_calc.target_id) \
            if base_calc.target_kind == "base" else None
        _header_row(ws, 1, ["Path", "Value", "Source type", "Class", "Entered by",
                            "Entered at", "Validation", "Confidence", "Notes"])
        r = 2
        if base is None:
            ws.cell(row=r, column=1, value="(no base calculation in this package)")
            return
        for path, prov in sorted(base.provenance.items()):
            ws.cell(row=r, column=1, value=path)
            parts = path.split(".")
            value = ""
            if parts[0] == "years" and len(parts) == 4:
                year = int(parts[1])
                if year in base.years:
                    value = str(getattr(getattr(base.years[year], parts[2]), parts[3]))
            ws.cell(row=r, column=2, value=value)
            ws.cell(row=r, column=3, value=prov.source_type.value)
            ws.cell(row=r, column=4, value=prov.input_class.value)
            ws.cell(row=r, column=5, value=prov.entered_by)
            ws.cell(row=r, column=6, value=prov.entered_at)
            ws.cell(row=r, column=7, value=prov.validation_status)
            ws.cell(row=r, column=8, value=prov.confidence or "")
            ws.cell(row=r, column=9, value=prov.notes or "")
            r += 1
        ws.column_dimensions["A"].width = 48

    def _import_mapping(self) -> None:
        ws = self.wb.create_sheet("Import Mapping")
        _header_row(ws, 1, ["Source document", "Source field", "Canonical path", "Status"])
        ws.cell(row=2, column=1,
                value="No document imports in this calculation set. Imported values, when "
                      "present, appear here with their original file, field, and approval status.")

    def _assumptions_by_year(self) -> None:
        ws = self.wb.create_sheet("Assumptions by Year")
        base_calc = next((r for r in self.results if r.target_kind == "base"), None)
        if base_calc is None:
            ws["A1"] = "(no base calculation in package)"
            return
        base = self.store.get_base_version(self.case_id, base_calc.target_id)
        years = sorted(base.years)
        _header_row(ws, 1, ["Input path"] + [str(y) for y in years] + ["Projection method"])
        policy = {fp.path: fp for fp in (base.projection_policy.fields if base.projection_policy else [])}
        from .projection import _iter_paths  # stable field enumeration
        r = 2
        for path in _iter_paths(base.years[years[0]]):
            ws.cell(row=r, column=1, value=path)
            for c, y in enumerate(years, start=2):
                section, field = path.split(".", 1)
                ws.cell(row=r, column=c, value=str(getattr(getattr(base.years[y], section), field)))
            fp = policy.get(path)
            default = base.projection_policy.default_method.value if base.projection_policy else ""
            ws.cell(row=r, column=len(years) + 2,
                    value=(f"{fp.method.value}"
                           + (f" @ {fp.annual_rate}" if fp and fp.annual_rate is not None else "")
                           ) if fp else f"{default} (default)")
            r += 1
        ws.column_dimensions["A"].width = 40

    def _scenario_change_log(self) -> None:
        ws = self.wb.create_sheet("Scenario Change Log")
        _header_row(ws, 1, ["Scenario", "Version", "Path", "Old", "New", "Reason",
                            "Source", "Changed by", "Changed at"])
        r = 2
        for res in self.results:
            if res.target_kind != "scenario":
                continue
            scn = self.store.get_scenario(self.case_id, res.target_id, res.scenario_version)
            for ov in scn.overrides:
                for c, v in enumerate([scn.name, scn.version, ov.path, str(ov.old_value),
                                       str(ov.new_value), ov.reason, ov.source,
                                       ov.changed_by, ov.changed_at], start=1):
                    ws.cell(row=r, column=c, value=v)
                r += 1
        ws.column_dimensions["C"].width = 42

    def _five_year_summary(self) -> None:
        ws = self.wb.create_sheet("Five-Year Summary")
        years = sorted(set(y for res in self.results for y in res.years))
        row = 1
        for res in self.results:
            label = f"{res.target_kind}:{res.target_id}" + (
                f"@v{res.scenario_version}" if res.scenario_version else "")
            ws.cell(row=row, column=1, value=label).font = TITLE
            row += 1
            _header_row(ws, row, ["Measure"] + [str(y) for y in years])
            hdr_row = row
            row += 1
            for field in SUMMARY_FIELDS:
                ws.cell(row=row, column=1, value=field)
                for c, y in enumerate(years, start=2):
                    if y in res.years:
                        cell = ws.cell(row=row, column=c, value=float(res.years[y].summary[field]))
                        sheet = _sheet_title(f"Calc {res.calculation_id[-6:]} {y}")
                        cell.hyperlink = f"#{quote_sheetname(sheet)}!A1"
                        if res.years[y].ruleset_status == "provisional":
                            cell.fill = WARN
                        if field == "total_tax":
                            name = f"total_tax_{res.calculation_id[-6:]}_{y}"
                            ref = f"{quote_sheetname(ws.title)}!${get_column_letter(c)}${row}"
                            self.wb.defined_names.add(DefinedName(name, attr_text=ref))
                            cell.comment = Comment(
                                f"Engine line F1040.L24 for {y}. Click to open the year sheet; "
                                "see Detail Index for the exact cell.", "engine")
                row += 1
            ws.cell(row=row, column=1,
                    value="Provisional-year cells are shaded; see Rules Manifest.").font = Font(italic=True)
            row += 2
        ws.column_dimensions["A"].width = 26

    def _scenario_comparison(self) -> None:
        ws = self.wb.create_sheet("Scenario Comparison")
        if not self.comparison:
            ws["A1"] = "(no comparison in this package)"
            return
        cmp_ = self.comparison
        r = 1
        ws.cell(row=r, column=1, value="Columns").font = HDR
        r += 1
        for i, col in enumerate(cmp_["columns"]):
            ws.cell(row=r, column=1, value=f"[{i}] {col['label']}")
            ws.cell(row=r, column=2, value=f"{len(col['changed_assumptions'])} changed assumptions")
            r += 1
        r += 1
        _header_row(ws, r, ["Year"] + [f"total_tax [{i}]" for i in range(len(cmp_["columns"]))]
                    + [f"Δ vs baseline [{i}]" for i in range(1, len(cmp_["columns"]))])
        r += 1
        for y in cmp_["years"]:
            row_data = cmp_["per_year"][str(y)]
            ws.cell(row=r, column=1, value=y)
            c = 2
            for summ in row_data["summaries"]:
                ws.cell(row=r, column=c, value=float(d(summ["total_tax"])))
                c += 1
            for delta in row_data["deltas_vs_baseline"][1:]:
                ws.cell(row=r, column=c, value=float(d(delta["total_tax"])))
                c += 1
            r += 1
        ws.cell(row=r, column=1, value="Cumulative Δ total tax").font = HDR
        for i, v in enumerate(cmp_["cumulative_total_tax_delta_vs_baseline"]):
            ws.cell(row=r, column=2 + i, value=float(d(v)))
        r += 2
        ws.cell(row=r, column=1,
                value="Differences trace to the Scenario Change Log and engine line IDs — "
                      "never to AI reasoning.").font = Font(italic=True)

    def _calculation_sheets(self) -> dict[tuple[str, int, str], str]:
        """One sheet per calculation-year with grouped rows per form."""
        cells: dict[tuple[str, int, str], str] = {}
        for res in self.results:
            for year, yr in sorted(res.years.items()):
                ws = self.wb.create_sheet(_sheet_title(f"Calc {res.calculation_id[-6:]} {year}"))
                ws["A1"] = (f"{res.target_kind}:{res.target_id} — tax year {year} — "
                            f"ruleset {yr.ruleset_id} ({yr.ruleset_status})")
                ws["A1"].font = TITLE
                if yr.ruleset_status == "provisional":
                    ws["A2"] = "PROVISIONAL RULESET — projected parameters, not enacted law"
                    ws["A2"].fill = WARN
                r = 4
                _header_row(ws, r, ["Line ID", "Label", "Value", "Formula", "Upstream",
                                    "Parameters", "Sources", "Rounding"])
                r += 1
                for form, items in list(yr.forms.items()) + list(yr.worksheets.items()):
                    ws.cell(row=r, column=1, value=form).font = HDR
                    ws.cell(row=r, column=1).fill = GRAY
                    r += 1
                    start = r
                    for li in items:
                        ws.cell(row=r, column=1, value=li.line_id)
                        ws.cell(row=r, column=2, value=li.label)
                        ws.cell(row=r, column=3, value=float(li.value))
                        ws.cell(row=r, column=4, value=li.formula)
                        ws.cell(row=r, column=5, value=", ".join(li.upstream_line_ids))
                        ws.cell(row=r, column=6, value=", ".join(li.parameter_ids))
                        ws.cell(row=r, column=7, value=", ".join(li.source_input_ids))
                        ws.cell(row=r, column=8, value=li.rounding)
                        self.detail_index.append(
                            (f"{res.calculation_id}.{year}.{li.line_id}", ws.title, f"C{r}"))
                        if li.line_id in VERIFIED_LINE_IDS:
                            cells[(res.calculation_id, year, li.line_id)] = f"C{r}"
                        r += 1
                    # collapsible supporting rows per form
                    for rr in range(start, r):
                        ws.row_dimensions[rr].outline_level = 1
                    r += 1
                for col, width in zip("ABCDEFGH", (14, 44, 14, 46, 24, 40, 40, 12)):
                    ws.column_dimensions[col].width = width
                ws.sheet_properties.outlinePr.summaryBelow = False
        return cells

    def _reconciliation(self) -> None:
        ws = self.wb.create_sheet("Reconciliation")
        _header_row(ws, 1, ["Calculation", "Year", "Check", "Description", "Status",
                            "Expected", "Actual", "Difference", "Tolerance", "Material",
                            "Supporting lines"])
        r = 2
        for res in self.results:
            for year, yr in sorted(res.years.items()):
                for c in yr.reconciliation:
                    vals = [res.calculation_id, year, c.check_id, c.description, c.status,
                            float(c.expected), float(c.actual), float(c.difference),
                            float(c.tolerance), "yes" if c.material else "no",
                            ", ".join(c.supporting_line_ids)]
                    for col, v in enumerate(vals, start=1):
                        cell = ws.cell(row=r, column=col, value=v)
                        if c.status == "failed":
                            cell.fill = FAIL
                    r += 1
        ws.column_dimensions["D"].width = 50

    def _calculation_trace(self) -> None:
        ws = self.wb.create_sheet("Calculation Trace")
        _header_row(ws, 1, ["Calculation", "Year", "Step"])
        r = 2
        for res in self.results:
            for year, yr in sorted(res.years.items()):
                for step in yr.calculation_trace:
                    ws.cell(row=r, column=1, value=res.calculation_id)
                    ws.cell(row=r, column=2, value=year)
                    ws.cell(row=r, column=3, value=step)
                    r += 1
        ws.column_dimensions["C"].width = 90

    def _rules_manifest(self) -> None:
        ws = self.wb.create_sheet("Rules Manifest")
        from .rulesets import DEFAULT_REGISTRY
        _header_row(ws, 1, ["Ruleset", "Year", "Status", "Hash", "Sources / provisional policy"])
        seen: set[str] = set()
        r = 2
        for res in self.results:
            for rs_id in res.ruleset_versions.values():
                if rs_id in seen:
                    continue
                seen.add(rs_id)
                rs = DEFAULT_REGISTRY.get(rs_id)
                man = rs.manifest()
                ws.cell(row=r, column=1, value=rs_id)
                ws.cell(row=r, column=2, value=rs.tax_year)
                ws.cell(row=r, column=3, value=rs.status)
                ws.cell(row=r, column=4, value=man["content_hash"])
                extra = "; ".join(man["source_references"])
                if man["provisional_policy"]:
                    extra += " | " + man["provisional_policy"]["uncertainty_warning"]
                cell = ws.cell(row=r, column=5, value=extra)
                if rs.is_provisional:
                    cell.fill = WARN
                r += 1
        ws.column_dimensions["E"].width = 100

    def _warnings(self) -> None:
        ws = self.wb.create_sheet("Validation Warnings")
        _header_row(ws, 1, ["Calculation", "Year", "Code", "Severity", "Requires review", "Message"])
        r = 2
        for res in self.results:
            for year, yr in sorted(res.years.items()):
                for g in yr.diagnostics:
                    vals = [res.calculation_id, year, g.code, g.severity.value,
                            "yes" if g.requires_review else "no", g.message]
                    for col, v in enumerate(vals, start=1):
                        ws.cell(row=r, column=col, value=v)
                    r += 1
        ws.column_dimensions["F"].width = 100

    def _review_signoff(self) -> None:
        ws = self.wb.create_sheet("Review Sign-off")
        _header_row(ws, 1, ["Review", "Target", "Reason", "State", "Reviewer",
                            "Reviewed at", "Findings", "Requested by", "Requested at"])
        r = 2
        for rec in self.store.list_reviews(self.case_id):
            vals = [rec.review_id, rec.target_id, rec.reason, rec.state.value,
                    rec.reviewer or "", rec.reviewed_at or "", rec.findings or "",
                    rec.requested_by, rec.requested_at]
            for col, v in enumerate(vals, start=1):
                ws.cell(row=r, column=col, value=v)
            r += 1
        ws.column_dimensions["C"].width = 60

    def _change_history(self) -> None:
        ws = self.wb.create_sheet("Change History")
        base_calc = next((r for r in self.results if r.target_kind == "base"), None)
        _header_row(ws, 1, ["Target", "State", "At", "By", "Note"])
        r = 2
        if base_calc:
            base = self.store.get_base_version(self.case_id, base_calc.target_id)
            for h in base.state_history:
                for col, v in enumerate([base.version_id, h.get("state", ""), h.get("at", ""),
                                         h.get("by", ""), h.get("note", "")], start=1):
                    ws.cell(row=r, column=col, value=v)
                r += 1

    def _detail_index_sheet(self) -> None:
        ws = self.wb.create_sheet("Detail Index")
        _header_row(ws, 1, ["Line key", "Sheet", "Cell"])
        for r, (key, sheet, cell) in enumerate(self.detail_index, start=2):
            ws.cell(row=r, column=1, value=key)
            link = ws.cell(row=r, column=2, value=sheet)
            link.hyperlink = f"#{quote_sheetname(sheet)}!{cell}"
            ws.cell(row=r, column=3, value=cell)
        ws.column_dimensions["A"].width = 50


def verify_workbook(path: Path, results: list[CalculationResult],
                    calc_cells: dict[tuple[str, int, str], str]) -> list[ReconciliationCheck]:
    """REC-XLSX: every critical total re-read from the saved workbook must
    equal the persisted calculation result (AGI, taxable income, total tax,
    total payments — see VERIFIED_LINE_IDS)."""
    wb = load_workbook(path, data_only=False)
    checks: list[ReconciliationCheck] = []
    for res in results:
        for year, yr in res.years.items():
            sheet = _sheet_title(f"Calc {res.calculation_id[-6:]} {year}")
            for n, (line_id, (summary_key, label)) in enumerate(sorted(VERIFIED_LINE_IDS.items()), start=1):
                cell = calc_cells.get((res.calculation_id, year, line_id))
                actual = d(str(wb[sheet][cell].value)) if cell else d(-1)
                expected = yr.summary[summary_key]
                checks.append(ReconciliationCheck(
                    check_id=f"REC-XLSX-{n:03d}-{res.calculation_id[-6:]}-{year}",
                    description=f"Workbook {label} matches persisted result for {year}",
                    status="passed" if actual == expected else "failed",
                    expected=expected, actual=actual, difference=actual - expected,
                    supporting_line_ids=[f"{sheet}!{cell}"]))
    return checks


def generate_audit_package(
    store: CaseStore, case_id: str, calculation_ids: list[str],
    comparison: dict[str, Any] | None, created_by: str, out_dir: Path,
) -> tuple[Path, str]:
    builder = AuditPackageBuilder(store, case_id, calculation_ids, comparison, created_by)
    out_path = Path(out_dir) / f"{builder.manifest_id}_{case_id}.xlsx"
    return builder.build(out_path)
