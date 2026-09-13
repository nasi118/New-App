"""Phase-8 workbook hardening tests: the immutable PackageManifest record,
the manifest reference + generation timestamp on the cover, four-way
critical-total verification, and visible failure on a tampered workbook."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ai_tax.excel_audit import VERIFIED_LINE_IDS, generate_audit_package, verify_workbook
from ai_tax.records import PackageManifest


@pytest.fixture()
def built_package(tmp_path, store, calc_service):
    from tests.conftest import make_approved_base

    case_id, base, result = make_approved_base(store, calc_service)
    path, file_hash = generate_audit_package(
        store, case_id, [result.calculation_id], comparison=None,
        created_by="tester", out_dir=tmp_path / "pkg")
    return store, case_id, result.calculation_id, path, file_hash


def test_manifest_record_stored_and_complete(built_package):
    store, case_id, calc_id, path, file_hash = built_package
    ids = store.backend.list_immutable("manifests")
    assert len(ids) == 1
    manifest = PackageManifest.model_validate(store.backend.get_immutable("manifests", ids[0]))
    assert manifest.case_id == case_id
    assert manifest.calculation_ids == [calc_id]
    assert manifest.file_sha256 == file_hash.removeprefix("sha256:")
    assert manifest.filename == path.name
    assert manifest.engine_version.startswith("engine-")
    assert manifest.ruleset_versions
    assert manifest.reconciliation_status == "passed"


def test_cover_names_manifest_and_timestamp(built_package):
    store, case_id, calc_id, path, _ = built_package
    wb = load_workbook(path)
    cover = wb["Cover"]
    text = " ".join(str(c.value) for row in cover.iter_rows() for c in row if c.value)
    manifest_id = store.backend.list_immutable("manifests")[0]
    assert manifest_id in text, "cover must reference the manifest id"
    assert "Generated at" in text


def test_verification_covers_all_critical_totals(built_package):
    store, case_id, calc_id, path, _ = built_package
    assert set(VERIFIED_LINE_IDS) == {"F1040.L11", "F1040.L15", "F1040.L24", "F1040.L33"}
    # Rebuild the cell map from the detail index sheet and verify clean.
    result = store.get_calculation(calc_id)
    wb = load_workbook(path)
    idx = wb["Detail Index"]
    cells = {}
    for row in idx.iter_rows(min_row=2, values_only=True):
        key, sheet, cell = row
        calc, year, line_id = key.rsplit(".", 2)[0], key.split(".")[-3], ".".join(key.split(".")[-2:])
        # key format: <calc_id>.<year>.<FORM>.<LINE>
        parts = key.split(".")
        calc = parts[0]
        year = int(parts[1])
        line_id = ".".join(parts[2:])
        if line_id in VERIFIED_LINE_IDS:
            cells[(calc, year, line_id)] = cell
    checks = verify_workbook(path, [result], cells)
    assert len(checks) == len(VERIFIED_LINE_IDS) * len(result.years)
    assert all(c.status == "passed" for c in checks)


def test_tampered_workbook_fails_verification(built_package, tmp_path):
    store, case_id, calc_id, path, _ = built_package
    result = store.get_calculation(calc_id)
    wb = load_workbook(path)
    idx = wb["Detail Index"]
    cells = {}
    target = None
    for row in idx.iter_rows(min_row=2, values_only=True):
        key, sheet, cell = row
        parts = key.split(".")
        line_id = ".".join(parts[2:])
        if line_id in VERIFIED_LINE_IDS:
            cells[(parts[0], int(parts[1]), line_id)] = cell
            if line_id == "F1040.L24" and target is None:
                target = (sheet, cell)
    sheet, cell = target
    wb[sheet][cell] = float(Decimal("1.00"))  # falsify total tax
    tampered = tmp_path / "tampered.xlsx"
    wb.save(tampered)
    checks = verify_workbook(tampered, [result], cells)
    failed = [c for c in checks if c.status == "failed"]
    assert failed and any("total tax" in c.description for c in failed)
