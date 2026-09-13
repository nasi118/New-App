"""Base-case lifecycle, scenario isolation, projections, comparison, workbook."""
from decimal import Decimal
from pathlib import Path

import pytest

from ai_tax.excel_audit import generate_audit_package
from ai_tax.projection import ProjectionError, project_years, rollforward_checks
from ai_tax.scenarios import ScenarioError, ScenarioService
from ai_tax.schemas import (
    CaseState,
    FieldProjection,
    ProjectionMethod,
    ProjectionPolicy,
    ReviewState,
)
from ai_tax.services import compare_scenarios
from ai_tax.store import StoreError

from .conftest import (
    default_policy,
    make_approved_base,
    mfj_household,
    single_household,
    some_provenance,
    year_inputs,
)


# -- lifecycle -------------------------------------------------------------

def test_illegal_transitions_rejected(store):
    case_id = store.create_case("u")
    base = store.create_base_version(case_id, single_household(), year_inputs(),
                                     default_policy(), some_provenance(), "u")
    with pytest.raises(StoreError, match="illegal"):
        store.transition_base(case_id, base.version_id, CaseState.APPROVED, "u")
    with pytest.raises(StoreError, match="illegal"):
        store.transition_base(case_id, base.version_id, CaseState.CALCULATED, "u")


def test_validation_failure_blocks_calculation(store, calc_service):
    case_id = store.create_case("u")
    bad = store.create_base_version(
        case_id, single_household(),
        year_inputs(income={"wages": "50000", "self_employment_income": "10000"}),
        default_policy(), some_provenance(), "u")
    version, diags = store.validate_base(case_id, bad.version_id, "u")
    assert version.state == CaseState.VALIDATION_FAILED
    assert any(g.code == "UNSUP-SE-001" for g in diags)
    with pytest.raises(StoreError, match="validate it before"):
        calc_service.run_for_base(case_id, bad.version_id, "u", "idem-x")


def test_approval_requires_review_state_and_supersedes(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    v2 = store.create_base_version(
        case_id, mfj_household(children=1),
        year_inputs(income={"wages": "190000"}), default_policy(),
        some_provenance(), "u", supersedes=base.version_id)
    with pytest.raises(StoreError, match="ready_for_review"):
        store.approve_base(case_id, v2.version_id, "reviewer")
    store.validate_base(case_id, v2.version_id, "u")
    calc_service.run_for_base(case_id, v2.version_id, "u", "idem-v2")
    store.transition_base(case_id, v2.version_id, CaseState.READY_FOR_REVIEW, "u")
    store.approve_base(case_id, v2.version_id, "reviewer")
    assert store.get_base_version(case_id, base.version_id).state == CaseState.SUPERSEDED
    assert store.get_base_version(case_id, v2.version_id).state == CaseState.APPROVED


def test_calculation_snapshots_write_once_and_tamper_evident(store, calc_service):
    case_id, base, result = make_approved_base(store, calc_service)
    with pytest.raises(StoreError, match="immutable"):
        store.record_calculation(result, "u")
    # tamper with the persisted snapshot -> integrity check must fail on read
    path = store.root / "calculations" / f"{result.calculation_id}.json"
    text = path.read_text().replace('"total_tax": "', '"total_tax": "9')
    path.write_text(text)
    with pytest.raises(StoreError, match="integrity"):
        store.get_calculation(result.calculation_id)


def test_review_resolution_immutable(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    rec = store.request_review(case_id, base.version_id, "materiality", [], "agent")
    store.resolve_review(case_id, rec.review_id, "reviewer", ReviewState.APPROVED, "ok")
    with pytest.raises(StoreError, match="already resolved"):
        store.resolve_review(case_id, rec.review_id, "reviewer", ReviewState.REJECTED, "no")


# -- projections -----------------------------------------------------------

def test_growth_rate_requires_explicit_rate():
    policy = ProjectionPolicy(base_year=2025, horizon_years=5, fields=[
        FieldProjection(path="income.wages", method=ProjectionMethod.GROWTH_RATE)])
    with pytest.raises(ProjectionError, match="annual_rate"):
        project_years(year_inputs(income={"wages": "100000"}), policy,
                      entered_by="u", entered_at="t")


def test_horizon_must_cover_five_years():
    policy = ProjectionPolicy(base_year=2025, horizon_years=3)
    with pytest.raises(ProjectionError, match="five years"):
        project_years(year_inputs(), policy, entered_by="u", entered_at="t")


def test_projection_growth_and_provenance():
    years, prov = project_years(
        year_inputs(income={"wages": "100000"}), default_policy(),
        entered_by="u", entered_at="t")
    assert sorted(years) == [2025, 2026, 2027, 2028, 2029]
    assert years[2026].income.wages == 103000
    assert years[2027].income.wages == 106090
    p = prov["years.2026.income.wages"]
    assert p.source_type.value == "projection"
    assert "growth_rate" in (p.notes or "")


def test_rollforward_reconciliation_detects_drift():
    base = year_inputs(income={"wages": "100000"})
    years, _ = project_years(base, default_policy(), entered_by="u", entered_at="t")
    ok = rollforward_checks(base, default_policy(), years)
    assert all(c.status == "passed" for c in ok)
    years[2027].income.wages = Decimal("999999")  # undeclared drift
    bad = rollforward_checks(base, default_policy(), years)
    assert any(c.status == "failed" and "years.2027.income.wages" in c.supporting_line_ids
               for c in bad)
    # the same change declared as an override is fine
    declared = rollforward_checks(base, default_policy(), years,
                                  override_paths={"years.2027.income.wages"})
    assert all(c.status == "passed" for c in declared)


# -- scenarios -------------------------------------------------------------

def test_scenario_requires_approved_base(store, calc_service):
    case_id = store.create_case("u")
    base = store.create_base_version(case_id, single_household(), year_inputs(),
                                     default_policy(), some_provenance(), "u")
    svc = ScenarioService(store)
    with pytest.raises(ScenarioError, match="APPROVED"):
        svc.create_scenario(case_id, base.version_id, "s", "r", "u")
    preview = svc.create_scenario(case_id, base.version_id, "s", "r", "u",
                                  allow_unapproved=True)
    assert preview.version == 1


def test_overrides_allowlisted_paths_only(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    svc = ScenarioService(store)
    scn = svc.create_scenario(case_id, base.version_id, "s", "r", "u")
    for bad_path in ["years.2026.household.filing_status",
                     "household.ctc_qualifying_children",
                     "years.2026.income.tax_year",
                     "years.2026.provenance.x"]:
        preview = svc.preview_overrides(
            scn, [{"path": bad_path, "new_value": "1", "reason": "nope"}])
        assert not preview["valid"], bad_path


def test_undeclared_old_value_rejected(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    svc = ScenarioService(store)
    scn = svc.create_scenario(case_id, base.version_id, "s", "r", "u")
    with pytest.raises(ScenarioError, match="undeclared mutation"):
        svc.apply_overrides(scn, [{"path": "years.2026.income.ira_conversion",
                                   "new_value": "50000", "old_value": "123",
                                   "reason": "conversion"}], "u", "idem-scn-bad")


def test_scenario_isolation_base_never_mutates(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    before_hash = store.get_base_version(case_id, base.version_id).input_snapshot_hash()
    svc = ScenarioService(store)
    scn = svc.create_scenario(case_id, base.version_id, "roth", "fill bracket", "u")
    scn2 = svc.apply_overrides(scn, [{"path": "years.2026.income.ira_conversion",
                                      "new_value": "75000", "old_value": "0",
                                      "reason": "Roth conversion"}], "u", "idem-scn-1")
    assert len(scn2.overrides) == 1 and scn2.version == 2
    after = store.get_base_version(case_id, base.version_id)
    assert after.input_snapshot_hash() == before_hash
    assert after.years[2026].income.ira_conversion == 0
    _, years, _ = svc.materialize(scn2)
    assert years[2026].income.ira_conversion == 75000


def test_apply_is_idempotent(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    svc = ScenarioService(store)
    scn = svc.create_scenario(case_id, base.version_id, "s", "r", "u")
    patch = [{"path": "years.2026.income.ira_conversion", "new_value": "10000",
              "old_value": "0", "reason": "x"}]
    a = svc.apply_overrides(scn, patch, "u", "idem-same-key")
    b = svc.apply_overrides(scn, patch, "u", "idem-same-key")
    assert a.version == b.version == 2
    assert len(store.list_scenarios(case_id)) == 2  # v1 + v2 only


def test_scenario_versions_immutable(store, calc_service):
    case_id, base, _ = make_approved_base(store, calc_service)
    svc = ScenarioService(store)
    scn = svc.create_scenario(case_id, base.version_id, "s", "r", "u")
    with pytest.raises(StoreError, match="immutable"):
        store.put_scenario(scn, "u")


# -- comparison + workbook -------------------------------------------------

def _base_and_scenario_calcs(store, calc_service):
    case_id, base, base_result = make_approved_base(store, calc_service)
    svc = ScenarioService(store)
    scn = svc.create_scenario(case_id, base.version_id, "roth", "fill bracket", "u")
    scn2 = svc.apply_overrides(scn, [{"path": "years.2026.income.ira_conversion",
                                      "new_value": "75000", "old_value": "0",
                                      "reason": "Roth conversion"}], "u", "idem-a")
    scn_result = calc_service.run_for_scenario(case_id, scn2.scenario_id, "u", "idem-b")
    return case_id, base_result, scn_result


def test_comparison_traces_to_overrides(store, calc_service):
    case_id, base_result, scn_result = _base_and_scenario_calcs(store, calc_service)
    cmp_ = compare_scenarios(store, [base_result.calculation_id, scn_result.calculation_id])
    assert cmp_["case_id"] == case_id
    assert cmp_["columns"][1]["changed_assumptions"][0]["path"] == \
        "years.2026.income.ira_conversion"
    delta_2026 = cmp_["per_year"]["2026"]["deltas_vs_baseline"][1]["total_tax"]
    assert Decimal(delta_2026) > 0
    assert all(c["status"] == "passed" for c in cmp_["delta_checks"])
    # no-op scenario with zero overrides but different results would fail REC-DELTA
    assert cmp_["all_reconciled"]


def test_cross_case_comparison_rejected(store, calc_service):
    _, r1, _ = _base_and_scenario_calcs(store, calc_service)
    case2, base2, r2 = make_approved_base(store, calc_service)
    with pytest.raises(ValueError, match="across cases"):
        compare_scenarios(store, [r1.calculation_id, r2.calculation_id])


def test_workbook_generated_verified_and_immutable(store, calc_service, tmp_path):
    case_id, base_result, scn_result = _base_and_scenario_calcs(store, calc_service)
    cmp_ = compare_scenarios(store, [base_result.calculation_id, scn_result.calculation_id])
    path, file_hash = generate_audit_package(
        store, case_id, [base_result.calculation_id, scn_result.calculation_id],
        cmp_, "u", tmp_path / "pkg")
    assert path.exists() and file_hash.startswith("sha256:")
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for sheet in ["Cover", "Case Metadata", "Base Inputs", "Import Mapping",
                  "Assumptions by Year", "Scenario Change Log", "Five-Year Summary",
                  "Scenario Comparison", "Reconciliation", "Calculation Trace",
                  "Rules Manifest", "Validation Warnings", "Review Sign-off",
                  "Change History", "Detail Index"]:
        assert sheet in wb.sheetnames, sheet
    # workbook total-tax cell equals persisted engine value (verified at build,
    # re-checked here)
    from ai_tax.excel_audit import _sheet_title
    ws = wb[_sheet_title(f"Calc {base_result.calculation_id[-6:]} 2025")]
    found = [row[2].value for row in ws.iter_rows()
             if row and row[0].value == "F1040.L24"]
    assert Decimal(str(found[0])) == base_result.years[2025].summary["total_tax"]
    # issued package is immutable: same filename cannot be registered again
    with pytest.raises(StoreError, match="immutable"):
        store.register_package(case_id, path.name, file_hash, "u")
