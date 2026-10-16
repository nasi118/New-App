"""Extract seed data from the three source workbooks into planner/src/seed.js.

The workbooks are not kept in git (scripts/check_no_generated_artifacts.py
rejects *.xlsx). To regenerate, copy them into planner/data/ and run:

    pip install openpyxl && python3 planner/extract_seed.py

Only raw inputs, labels, notes and source citations are pulled from the
workbooks. Every computed figure is recomputed by the app engines.
"""
import json, re, glob, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
def wbpath(part):
    return glob.glob(os.path.join(HERE, 'data', f'*{part}*.xlsx'))[0]

def sources(ws, first=4):
    out = []
    for r in range(first, ws.max_row + 1):
        sid = ws.cell(r, 1).value
        if not sid:
            continue
        out.append({'id': sid, 'title': ws.cell(r, 2).value, 'ref': ws.cell(r, 3).value,
                    'excerpt': ws.cell(r, 4).value})
    return out

# ---------------------------------------------------------------- HNWI
wb = openpyxl.load_workbook(wbpath('hnwi'))
cov = wb['Cover & Assumptions']
assumptions = []
for r in range(10, 29):
    name = cov.cell(r, 1).value
    if name:
        assumptions.append({'name': name, 'desc': cov.cell(r, 2).value, 'value': cov.cell(r, 3).value,
                            'format': cov.cell(r, 4).value, 'source': cov.cell(r, 5).value})

def norm_formula(f, row):
    """Excel row formula -> row-relative expression (G..N refer to this row)."""
    f = f.lstrip('=')
    f = re.sub(r'([A-Z])%d\b' % row, r'\1', f)
    f = re.sub(r'(\d+(?:\.\d+)?)%', lambda m: str(float(m.group(1)) / 100), f)
    return f

sc = wb['Strategy Comparison']
strategies = []
for r in range(4, 54):
    g = lambda c: sc.cell(r, c).value
    strategies.append({
        'no': g(1), 'title': g(2), 'situation': g(3), 'objective': g(4), 'mechanics': g(5),
        'category': g(6), 'G': g(7), 'H': g(8), 'I': g(9), 'J': g(10), 'K': g(11), 'L': g(12),
        'formulaDesc': g(13), 'formula': norm_formula(g(14), r), 'low': norm_formula(g(15), r),
        'high': norm_formula(g(16), r), 'risk': g(17), 'implementation': g(18),
        'jurisdiction': g(19), 'sources': g(20),
    })

cb = wb['Cost-Benefit Summary']
for r in range(15, 65):
    no = cb.cell(r, 2).value
    s = next(s for s in strategies if s['no'] == no)
    s['selected'] = cb.cell(r, 1).value == 'Yes'
    s['conflict'] = cb.cell(r, 14).value

ic = wb['Implementation Checklist']
checklist = []
for r in range(12, 212):
    g = lambda c: ic.cell(r, c).value
    checklist.append({'id': f'IC-{r-11:03d}', 'no': g(1), 'phase': g(3), 'item': g(4), 'party': g(5),
                      'docs': g(6), 'timeline': g(7), 'cost': g(8), 'deps': g(9), 'notes': g(10),
                      'status': g(11), 'priority': g(12)})

hnwi = {
    'client': cov['B3'].value, 'prepared': cov['B4'].value, 'purpose': cov['B5'].value,
    'note': cov['B6'].value, 'horizon': cb['C4'].value, 'assumptions': assumptions,
    'strategies': strategies, 'checklist': checklist, 'sources': sources(wb['Sources']),
}

# ---------------------------------------------------------------- Rental
wb = openpyxl.load_workbook(wbpath('rental'))
ai = wb['Assumptions_Input']
es = wb['Entity Structure']
agg = wb['199A_Aggregation']
props = []
for r in range(14, 18):
    g = lambda c: ai.cell(r, c).value
    props.append({'name': g(1), 'entity': g(2), 'type': g(3), 'location': g(4), 'acquired': g(5),
                  'rents': g(6), 'otherIncome': g(7), 'opex': g(8), 'repairs': g(9), 'taxIns': g(10),
                  'interest': g(11), 'depreciation': g(12), 'wages': g(13), 'ubia': g(14),
                  'hours': g(15), 'notes': g(16)})
rental = {
    'purpose': wb['README']['B4'].value,
    'owners': [{'name': es.cell(r, 1).value, 'pct': es.cell(r, 4).value, 'note': es.cell(r, 7).value}
               for r in (4, 5)],
    'holdco': {'name': es['A6'].value, 'status': es['B6'].value},
    'hours': {'johnRP': ai['B4'].value, 'johnTotal': ai['B5'].value, 'janeRP': ai['B6'].value,
              'combinedRental': ai['B7'].value, 'outsideMgmt': ai['B8'].value},
    'taxableIncome': ai['B9'].value, 'netCapitalGain': ai['B10'].value,
    'properties': props,
    'aggregation': {'majorityYear': agg['B5'].value, 'sameYear': agg['B6'].value,
                    'noSSTB': agg['B7'].value, 'factorA': agg['B8'].value,
                    'factorB': agg['B9'].value, 'factorC': agg['B10'].value},
    'aggregationNotes': {k: agg.cell(r, 6).value for k, r in
                         [('common', 4), ('majorityYear', 5), ('sameYear', 6), ('noSSTB', 7),
                          ('factorA', 8), ('factorB', 9), ('factorC', 10)]},
    'sources': sources(wb['Sources']),
}

# ---------------------------------------------------------------- 1040 workpaper
wb = openpyxl.load_workbook(wbpath('1040'))
sd = wb['Source_Docs']
def rows(ws, r0, r1, keys, c0=1):
    return [{k: ws.cell(r, c0 + i).value for i, k in enumerate(keys)} for r in range(r0, r1 + 1)]
w2 = rows(sd, 5, 6, ['who', 'employer', 'wages', 'fedwh', 'ssWages', 'ssTax', 'medWages', 'medTax', 'stateWages', 'stateWh'])
ints = rows(sd, 12, 14, ['payer', 'taxable', 'exempt', 'wh', 'note'])
divs = rows(sd, 20, 21, ['payer', 'ordinary', 'qualified', 'capGainDist', 'wh', 'note'])
sales = rows(sd, 27, 28, ['desc', 'term', 'proceeds', 'basis', 'adj'])
other = {k: sd[c].value for k, c in [('stateRefund', 'B34'), ('unemployment', 'B35'), ('gambling', 'B36'),
         ('mortgage', 'B37'), ('stateIncomeTax', 'B38'), ('realEstateTax', 'B39'), ('personalPropTax', 'B40'),
         ('cashCharity', 'B41'), ('estPayments', 'B42'), ('extensionPayment', 'B43'), ('educationCredit', 'B44')]}
scc = wb['Sch_C_Business']
schc_exp = [{'line': scc.cell(r, 1).value, 'desc': scc.cell(r, 2).value, 'amount': scc.cell(r, 3).value,
             'note': scc.cell(r, 8).value} for r in range(11, 20)]
se_ = wb['Sch_E_Rental_Royalty']
ecols = ['gross', 'advertising', 'auto', 'cleaning', 'commissions', 'insurance', 'mortgage', 'taxes', 'repairs', 'depreciation', 'other']
schE = []
for r in (4, 5):
    d = {'name': se_.cell(r, 1).value, 'type': se_.cell(r, 2).value}
    d.update({k: se_.cell(r, 3 + i).value for i, k in enumerate(ecols)})
    schE.append(d)
k1 = wb['K1_Input']
k1keys = ['entity', 'form', 'ein', 'ptp', 'passive', 'material', 'box1', 'box2', 'interest', 'ordDiv', 'qualDiv',
          'stcg', 'ltcg', 's179', 'charity', 'ftc', 'qbi', 'se', 'distributions', 'basis']
k1rows = rows(k1, 4, 7, k1keys)
for d in k1rows:
    for f in ('ptp', 'passive', 'material'):
        d[f] = d[f] == 'Yes'
dep = wb['Schedule_8812_CTC']
deps = [{'name': dep.cell(r, 1).value, 'age': dep.cell(r, 2).value, 'ssn': dep.cell(r, 3).value == 'Yes',
         'ctc': dep.cell(r, 4).value == 'Yes', 'odc': dep.cell(r, 5).value == 'Yes', 'note': dep.cell(r, 8).value}
        for r in (4, 5)]
oi = wb['Open_Items']
open_items = [{'area': oi.cell(r, 2).value, 'entity': oi.cell(r, 3).value, 'issue': oi.cell(r, 4).value,
               'amount': oi.cell(r, 5).value, 'priority': oi.cell(r, 6).value, 'owner': oi.cell(r, 7).value,
               'status': oi.cell(r, 10).value, 'resolution': ''} for r in range(4, 7)]
tc = wb['Training_Checklist']
training = [{'no': tc.cell(r, 1).value, 'area': tc.cell(r, 2).value, 'item': tc.cell(r, 3).value,
             'prepared': False, 'reviewed': False} for r in range(4, 21)]
steps = [{'no': s.cell(r, 1).value, 'area': s.cell(r, 2).value, 'action': s.cell(r, 3).value,
          'status': s.cell(r, 6).value} for s in [wb['Start_Here']] for r in range(10, 18)]
prof = wb['Demo_Client_Profile']
w1040 = {
    'profile': {prof.cell(r, 1).value: prof.cell(r, 2).value for r in range(3, 13)},
    'tolerance': wb['Start_Here']['B5'].value,
    'steps': steps,
    'w2': w2, 'interest': ints, 'dividends': divs, 'sales': sales, 'other': other,
    'schA': {'medical': wb['Sch_A_Itemized']['C4'].value, 'casualty': 0, 'otherItemized': 0,
             'standard': wb['Sch_A_Itemized']['C18'].value, 'saltCap': 40000},
    'schC': {'gross': scc['C4'].value, 'returns': scc['C5'].value, 'cogs': scc['C7'].value,
             'otherIncome': scc['C9'].value, 'expenses': schc_exp, 'homeOffice': scc['C22'].value},
    'schE': schE, 'k1': k1rows,
    'pal': {'rreIncome': 0, 'priorUnallowed': 0, 'threshold': 150000,
            'magi': wb['Form_8582_PAL']['C11'].value, 'passiveIncome': 0},
    'se': {'factor': 0.9235, 'wageBase': wb['Schedule_SE']['C9'].value, 'owner': 'Casey'},
    'qbi': {'reitPtp': 0},
    'deps': deps, 'ctcAmount': 2200, 'odcAmount': 500, 'actcMax': 1700,
    'sch1': {'alimony': 0, 'otherGains': 0, 'farm': 0, 'hsa': 3000, 'studentLoan': 1200, 'otherAdj': 0},
    'sch2': {'otherAdditions': 0, 'amt': 0, 'addlMedicare': 0, 'niit': wb['Schedule_2']['C9'].value, 'otherTaxes': 0},
    'sch3': {'childCare': 0, 'retirementSavers': 0, 'energy': 0, 'otherNonref': 0, 'ptc': 0,
             'excessSS': 0, 'fuel': 0, 'otherRefundable': 0},
    'taxCalc': {'softwareTax': wb['Tax_Calc']['B7'].value, 'method': wb['Tax_Calc']['B8'].value,
                'penalty': wb['Tax_Calc']['B9'].value},
    'misc': {'l1b': 0, 'l1c': 0, 'l1d': 0, 'l1e': 0, 'l1f': 0, 'l1g': 0, 'l1h': 0, 'l4b': 0, 'l5b': 0,
             'l6b': 0, 'l13b': 0, 'l25c': 0, 'l27a': 0, 'l29': 0, 'l30': 0, 'l36': 0},
    'openItems': open_items, 'training': training,
    'signoff': {'preparedBy': '', 'preparedDate': '', 'reviewedBy': '', 'reviewDate': '', 'notes': ''},
    'sources': sources(wb['Sources']),
}

seed = {'hnwi': hnwi, 'rental': rental, 'w1040': w1040}
with open(os.path.join(HERE, 'src', 'seed.js'), 'w') as f:
    f.write('/* Generated by extract_seed.py from the three source workbooks. Do not edit by hand. */\n')
    f.write('window.SEED = ' + json.dumps(seed, ensure_ascii=False, default=str) + ';\n')
print('strategies', len(strategies), 'checklist', len(checklist), 'bytes', os.path.getsize(os.path.join(HERE, 'src', 'seed.js')))
